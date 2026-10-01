"""One-off: add the 'Role settings' sheet (and a how-to section) to the rules workbook."""
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from rules_io import RULES_FILE

ANCILLARY = "garage; garages; shed; carport; roof; hut; service; outbuilding; greenhouse; barn; container; shelter; toilets; kiosk; stable; farm_auxiliary; storage_tank; silo; static_caravan; cabin"
PRIMARY = "school; kindergarten; college; university; hospital; healthcare; hangar; government; public; civic; fire_station; train_station; terminal"

wb = openpyxl.load_workbook(RULES_FILE)
if "Role settings" in wb.sheetnames:
    del wb["Role settings"]
ws = wb.create_sheet("Role settings", 2)
head = PatternFill("solid", fgColor="1F3864")
ws.append(["Setting", "Value", "Explanation"])
ws.append(["Minimum primary footprint (m²)", 150,
           "A building at a multi-building site is PRIMARY if its footprint is at least this big. Smaller = ANCILLARY."])
ws.append(["Ancillary OSM building values", ANCILLARY,
           "If OSM tags a building with one of these (building=garage etc.) it is ANCILLARY whatever its size. Separate with ;"])
ws.append(["Primary OSM building values", PRIMARY,
           "If OSM tags a building with one of these it is PRIMARY whatever its size. Separate with ;"])
ws.append([])
ws.append(["Building Type", "Minimum primary footprint (m²)",
           "Optional: set a different minimum for one building type. Leave blank to use the default above."])
types = []
for r in wb["Rules"].iter_rows(min_row=2, values_only=True):
    if r[4] and r[4] not in types and r[3]:
        types.append(r[4])
for t in types:
    ws.append([t, None, ""])
for row in (1, 6):
    for c in ws[row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
for col, w in zip("ABC", [42, 60, 100]):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2, max_row=4):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")

h = wb["How to edit"]
for text, style in [
    ("", None),
    ("Building roles (Role settings sheet)", "h"),
    ("Each building at a site is labelled Primary, Ancillary, Standalone (only building at its site), Whole site or Unknown.", None),
    ("Change the minimum footprint, or the OSM building values, on the 'Role settings' sheet. Optional per-type minimums are below them.", None),
    ("Individual buildings can be corrected in outputs/overrides.csv using the 'Building Role Override' column.", None),
]:
    h.append([text])
    c = h.cell(row=h.max_row, column=1)
    if style == "h":
        c.font = Font(bold=True, size=12)
wb["Change log"].append(["Role settings", "New sheet controlling the Building Role flag (thresholds and OSM building values)"])
wb.save(RULES_FILE)
print("added Role settings with", len(types), "building types")
