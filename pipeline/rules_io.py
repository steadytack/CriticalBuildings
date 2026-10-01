"""Load and check the Rules sheet. Used by every pipeline step.

Run directly (`python pipeline/rules_io.py`) to check the rules and get plain-English messages.
"""
import difflib
import hashlib
import re
import sys
from dataclasses import dataclass
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
RULES_FILE = ROOT / "Critical Building Data Product (tidied).xlsx"

COLUMNS = ["Rule Ref", "Active (Y/N)", "Priority", "Primary Category", "Building Type",
           "Operational Importance", "Operational Importance Driver", "Source", "Tag", "Notes"]

# Tag keys we expect to see in OSM. A key outside this list gets a "did you mean?" hint.
KNOWN_OSM_KEYS = {
    "amenity", "healthcare", "emergency", "social_facility", "social_facility:for", "building",
    "school", "railway", "public_transport", "station", "landuse", "aeroway", "government",
    "military_service", "leisure", "telecom", "plant:output:hot_water", "plant:output:electricity",
    "military", "office", "man_made", "power", "tourism", "shop", "healthcare:speciality",
    "telecom:data_center",
}
OS_GROUPS = {"Emergency services", "Education", "Transport", "Government Buildings",
             "Miscellaneous", "Medical Care", "Utilities", "Sport And Entertainment",
             # OS Open Map Local themes
             "Air Transport", "Attraction And Leisure", "Cultural Facility", "Emergency Service",
             "Religious Buildings", "Retail", "Road Transport", "Sports Or Exercise Facility",
             "Water Transport", "Rail Transport"}
IMPORTANCE = {"High", "Medium", "Low"}


@dataclass
class Tag:
    rule_ref: int
    priority: int
    category: str
    building_type: str
    importance: str
    driver: str
    source: str      # OSM / OS / MANUAL
    key: str         # OSM key, or OS group
    value: str       # OSM value ('*' = any), or OS classification
    raw: str
    notes: str = ""


def _rows(path=RULES_FILE):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Rules"]
    header = [c.value for c in ws[1]]
    return header, [(i, r) for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2)
                    if any(v not in (None, "") for v in r)]


def check_rules(path=RULES_FILE):
    """Return (errors, warnings, tags). Messages are plain English naming rule and row."""
    errors, warnings, tags = [], [], []
    header, rows = _rows(path)
    if header[:len(COLUMNS)] != COLUMNS:
        errors.append(f"Rules sheet headings are not as expected. Expected: {', '.join(COLUMNS)}")
        return errors, warnings, tags
    types_by_ref = {}
    for rownum, r in rows:
        ref, active, prio, cat, btype, imp, driver, source, tag = r[:9]
        where = f"Rule {ref} (row {rownum})"
        if not isinstance(ref, (int, float)):
            errors.append(f"Row {rownum}: Rule Ref must be a number, found '{ref}'.")
            continue
        ref = int(ref)
        where = f"Rule {ref} (row {rownum})"
        if str(active).strip().upper() not in ("Y", "N"):
            errors.append(f"{where}: 'Active' must be Y or N, found '{active}'.")
            continue
        if str(active).strip().upper() == "N":
            continue
        if not isinstance(prio, (int, float)):
            errors.append(f"{where}: Priority must be a number (lower wins), found '{prio}'.")
            continue
        if str(imp).strip() not in IMPORTANCE:
            errors.append(f"{where}: Operational Importance must be High, Medium or Low, found '{imp}'.")
            continue
        for label, val in (("Primary Category", cat), ("Building Type", btype), ("Driver", driver)):
            if val in (None, ""):
                errors.append(f"{where}: '{label}' is empty.")
        types_by_ref.setdefault(ref, set()).add((cat, btype, imp))
        source = str(source).strip().upper() if source else ""
        tag = (str(tag).strip() if tag is not None else "")
        if source == "MANUAL":
            warnings.append(f"{where}: '{btype}' has no automatic tag - it only appears via manual override.")
            continue
        if source not in ("OSM", "OS"):
            errors.append(f"{where}: Source must be OSM, OS or MANUAL, found '{source}'.")
            continue
        if not tag:
            errors.append(f"{where}: Tag is empty.")
            continue
        if source == "OSM":
            if not re.fullmatch(r"[a-z_:0-9]+=(\*|[A-Za-z0-9_:;.\- ]+)", tag) or " " in tag.split("=")[0]:
                errors.append(f"{where}: OSM tag '{tag}' should look like key=value with no spaces "
                              f"(e.g. amenity=fire_station).")
                continue
            key, value = tag.split("=", 1)
            if key not in KNOWN_OSM_KEYS:
                hint = difflib.get_close_matches(key, KNOWN_OSM_KEYS, n=1)
                warnings.append(f"{where}: OSM key '{key}' is not one I recognise"
                                + (f" - did you mean '{hint[0]}'?" if hint else "."))
        else:
            if tag.count(">") != 1:
                errors.append(f"{where}: OS tag '{tag}' should look like 'Group > Classification'.")
                continue
            key, value = [s.strip() for s in tag.split(">")]
            if key not in OS_GROUPS:
                hint = difflib.get_close_matches(key, OS_GROUPS, n=1)
                warnings.append(f"{where}: OS group '{key}' is not one I recognise"
                                + (f" - did you mean '{hint[0]}'?" if hint else "."))
        tags.append(Tag(ref, int(prio), str(cat).strip(), str(btype).strip(), str(imp).strip(),
                        str(driver).strip(), source, key, value, tag, str(r[9] or "").strip()))
    for ref, variants in types_by_ref.items():
        if len(variants) > 1:
            errors.append(f"Rule {ref}: its rows disagree on Category/Building Type/Importance "
                          f"({len(variants)} different versions). Rows of one rule must match.")
    return errors, warnings, tags


def rule_version(path=RULES_FILE):
    """Stable version stamp: date of last save + short hash of the Rules sheet content."""
    _, rows = _rows(path)
    h = hashlib.sha1(repr([r for _, r in rows]).encode()).hexdigest()[:6]
    return f"{Path(path).stat().st_mtime_ns and __import__('datetime').date.fromtimestamp(Path(path).stat().st_mtime)}-{h}"


def load_rules(path=RULES_FILE):
    errors, warnings, tags = check_rules(path)
    if errors:
        raise SystemExit("Rules check FAILED:\n  " + "\n  ".join(errors))
    return tags, warnings

DEFAULT_ROLE_SETTINGS = dict(min_area=150.0, ancillary=set(), primary=set(), per_type={})


def load_role_settings(path=RULES_FILE):
    """Read the 'Role settings' sheet. Returns (settings, errors)."""
    errors = []
    wb = openpyxl.load_workbook(path, data_only=True)
    if "Role settings" not in wb.sheetnames:
        return dict(DEFAULT_ROLE_SETTINGS), ["The 'Role settings' sheet is missing."]
    ws = wb["Role settings"]
    s = dict(min_area=150.0, ancillary=set(), primary=set(), per_type={})
    in_types = False
    for rownum, (a, b, _c) in enumerate(ws.iter_rows(min_row=2, max_col=3, values_only=True), start=2):
        if a == "Building Type":
            in_types = True
            continue
        if a in (None, ""):
            continue
        if not in_types:
            if a.startswith("Minimum primary"):
                try:
                    s["min_area"] = float(b)
                except (TypeError, ValueError):
                    errors.append(f"Role settings row {rownum}: minimum footprint must be a number, found '{b}'.")
            elif a.startswith("Ancillary"):
                s["ancillary"] = {v.strip() for v in str(b or "").split(";") if v.strip()}
            elif a.startswith("Primary"):
                s["primary"] = {v.strip() for v in str(b or "").split(";") if v.strip()}
        elif b not in (None, ""):
            try:
                s["per_type"][a] = float(b)
            except (TypeError, ValueError):
                errors.append(f"Role settings row {rownum}: minimum for '{a}' must be a number, found '{b}'.")
    both = s["ancillary"] & s["primary"]
    if both:
        errors.append(f"Role settings: {', '.join(sorted(both))} appear in both the ancillary and primary lists.")
    return s, errors


def load_reference_mapping(path=RULES_FILE):
    """Read the 'Reference mapping' sheet. Returns (rows, errors); rows are dicts for ACTIVE rows only."""
    errors, rows = [], []
    wb = openpyxl.load_workbook(path, data_only=True)
    if "Reference mapping" not in wb.sheetnames:
        return rows, ["The 'Reference mapping' sheet is missing."]
    valid_types = {str(r[4]).strip() for r in wb["Rules"].iter_rows(min_row=2, values_only=True) if r[4]}
    for n, r in enumerate(wb["Reference mapping"].iter_rows(min_row=2, max_col=7, values_only=True), start=2):
        if not any(r):
            continue
        src, field, match, value, btype, active = [(str(x).strip() if x is not None else "") for x in r[:6]]
        where = f"Reference mapping row {n}"
        if active.upper() != "Y":
            continue
        if src not in ("DfE", "CQC", "NHS ODS"):
            errors.append(f"{where}: Reference Source must be DfE, CQC or NHS ODS, found '{src}'.")
        if match not in ("Contains", "Equals"):
            errors.append(f"{where}: Match must be Contains or Equals, found '{match}'.")
        if not field or not value:
            errors.append(f"{where}: Field and Value must be filled in.")
        if btype != "IGNORE" and btype not in valid_types:
            errors.append(f"{where}: Building Type '{btype}' is not a Building Type on the Rules sheet.")
        rows.append(dict(source=src, field=field, match=match, value=value, building_type=btype))
    return rows, errors


if __name__ == "__main__":
    errs, warns, tags = check_rules()
    errs += load_role_settings()[1]
    errs += load_reference_mapping()[1]
    print(f"Rules file: {RULES_FILE.name}   version: {rule_version()}")
    for w in warns:
        print("  note:", w)
    for e in errs:
        print("  ERROR:", e)
    print(f"{len(tags)} active tag rows across {len({t.rule_ref for t in tags})} rules; "
          f"{len(errs)} errors, {len(warns)} notes")
    sys.exit(1 if errs else 0)
