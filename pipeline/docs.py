"""Data dictionary, methodology, licences and release notes - generated on every run so they never go stale.

Writes outputs/critical_buildings_data_dictionary.xlsx, outputs/ATTRIBUTION.txt and
outputs/critical_buildings_summary.csv (a slim, easy view of the main dataset).
"""
import datetime as dt
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

import merge
import refmatch
import sites as sitelib

# ------------------------------------------------------------------ data quality score
VERIFICATION_POINTS = {"Verified": 40, "Verified (manual)": 40, "Authoritative only": 40, "Not covered": 20,
                       "Unverified": 10, "Closed per register": 0}
SOURCE_POINTS = {"Both": 25, "OSM": 12, "OS Important Buildings": 12}   # other sources are registers
REGISTER_POINTS = 20
CONFIDENCE_POINTS = {"High": 20, "Medium": 12, "Low": 5}
POLYGON_POINTS, POINT_POINTS = 15, 5
OPEN_REVIEW_PENALTY = -10
BANDS = "High 75-100, Medium 50-74, Low 0-49"


def quality_score(df):
    """0-100 score from how well a record is evidenced. A heuristic to help sorting and filtering, not a probability."""
    ver = df["Verification"].map(VERIFICATION_POINTS).fillna(10)
    src = df["Source"].map(SOURCE_POINTS).fillna(REGISTER_POINTS)
    conf = df["Confidence"].map(CONFIDENCE_POINTS).fillna(12)
    geom = np.where(df["Geom"].str.upper().str.startswith("POINT"), POINT_POINTS, POLYGON_POINTS)
    pen = np.where(df["Review Status"] == "Open", OPEN_REVIEW_PENALTY, 0)
    return (ver + src + conf + geom + pen).clip(0, 100).astype(int)


# ------------------------------------------------------------------ column dictionary
# name: (group, description, values / example)
D = {
    # identity
    "Building ID": ("Identity", "Stable unique ID for the building record. Kept across runs when the same real building is found again.", "BLDG-000123"),
    "Source": ("Identity", "Which dataset(s) the record comes from.", "OSM; OS Important Buildings; Both; DfE GIAS; CQC; NHS ODS"),
    "Source Feature ID": ("Identity", "ID(s) of the source feature(s) so the record can be traced back. 'Both' shows OS id / OSM id.", "way/7000000; <OS GUID> / way/123; CQC:1-12345"),
    "Name": ("Identity", "Best available name of the building.", "Hartismere Hospital"),
    # classification
    "Primary Category": ("Classification", "Final category after any override: Emergency, Care, Education, Transport, Government Buildings, Miscellaneous.", "Education"),
    "Building Type": ("Classification", "Final building type after any override (a type from the Rules sheet).", "Primary School"),
    "Operational Importance": ("Classification", "Final operational importance after any override.", "High; Medium; Low"),
    "Primary Category Default": ("Classification", "Category assigned by the rules (before overrides).", "Emergency"),
    "Building Type Default": ("Classification", "Building type assigned by the rules (before overrides).", "Fire Station"),
    "Operational Importance Default": ("Classification", "Importance assigned by the rules.", "High"),
    "Operational Importance Driver Default": ("Classification", "Plain-English reason for the importance, from the rules.", "Emergency response"),
    "Rule Notes": ("Classification", "Notes on the rule that classified the record, from the Rules sheet.", ""),
    "Classification Note": ("Classification", "Set when an authoritative source changed the building type chosen from OS/OSM.", "Building Type set from CQC (...); OS/OSM said Hospital"),
    "OS Rule Ref": ("Classification", "Rule number matched from OS Important Buildings (blank if none).", "17"),
    "OSM Rule Ref": ("Classification", "Rule number matched from OSM tags (blank if none).", "18"),
    "All Matched Rule Refs": ("Classification", "Every rule number the record matched, separated by ;", "17;20"),
    # overrides
    "Primary Category Override": ("Overrides", "Manual category; when filled it replaces the default.", ""),
    "Building Type Override": ("Overrides", "Manual building type; replaces the default.", ""),
    "Operational Importance Override": ("Overrides", "Manual importance; replaces the default.", ""),
    "Operational Importance Driver Override": ("Overrides", "Manual reason for the importance.", ""),
    "Manual Override Notes": ("Overrides", "Free-text reason for any manual decision on this record.", ""),
    "Building Role Override": ("Overrides", "Manual building role; replaces the default role.", ""),
    "Review Decision": ("Overrides", "Decision recorded against a review item (for example Accepted).", "Accepted"),
    "Verification Override": ("Overrides", "Manual verification result for the whole site.", "Verified (manual)"),
    "Excluded": ("Overrides", "True if a manual decision removed the record. Excluded records are not in the delivered files.", "False"),
    # site
    "Site ID": ("Site", "Stable ID of the site (school, hospital, airfield ...) the building belongs to.", "SITE-000012"),
    "Site Name": ("Site", "Name of the site.", "Culford School"),
    "Site Type": ("Site", "Most common building type among the site's buildings.", "Primary School"),
    "Site Method": ("Site", "How the site was formed: OS Functional Site, OSM site, OSM area, Proximity cluster, Single building, Authoritative source.", "OS Functional Site"),
    "Buildings in Site": ("Site", "Number of building records in the site.", "12"),
    # role
    "Building Role": ("Building role", "Final role of the building within its site after any override.", "Primary; Ancillary; Standalone; Whole site; Unknown"),
    "Building Role Default": ("Building role", "Role chosen by the pipeline.", "Primary"),
    "Building Role Reason": ("Building role", "Plain-English reason for the role.", "Footprint 420 m2 is at least 150 m2"),
    # authoritative
    "Verification": ("Authoritative check", "Whether DfE / CQC / NHS ODS confirms the site. Verified, Verified (manual), Authoritative only (added from the register), Unverified (type is covered but nothing matched), Closed per register, Not covered (no register for this type).", "Verified"),
    "Verification Note": ("Authoritative check", "Why a site is unverified or how far a match was.", "No authoritative record found within 300 m"),
    "Authoritative Distance (m)": ("Authoritative check", "Distance in metres between the site and the matched authoritative point. Over 500 is sent for review.", "35"),
    "Authoritative Source": ("Authoritative check", "Register(s) that matched the site.", "DfE GIAS; CQC; NHS ODS"),
    "Authoritative ID": ("Authoritative check", "Register ID(s): DfE URN, CQC Location ID, NHS ODS code.", "124807"),
    "Authoritative Name": ("Authoritative check", "Name(s) the register uses.", ""),
    "Authoritative Status": ("Authoritative check", "Status from the register.", "Open; Active"),
    "Capacity": ("Authoritative check", "Pupils (DfE) or beds (CQC), summed over matched records.", "243"),
    "Capacity Unit": ("Authoritative check", "Unit of Capacity.", "pupils; beds"),
    "Authoritative Detail": ("Authoritative check", "Type/phase (DfE) or inspection category (CQC).", "Academy converter / Primary"),
    "CQC Rating": ("Authoritative check", "Latest overall CQC rating.", "Good"),
    "Operator / Trust": ("Authoritative check", "Trust, provider or operator.", ""),
    "Authoritative Postcode": ("Authoritative check", "Postcode from the register.", ""),
    "UPRN": ("Authoritative check", "UPRN(s) given by CQC for the location. Not a full UPRN count - see Count UPRNs.", ""),
    "ODS Code": ("Authoritative check", "NHS ODS organisation code.", ""),
    # quality and review
    "Match Method": ("Quality and review", "How the record was built: Overlap, Proximity, Proximity (names agree), OS only, OSM only, OSM area (no buildings), Authoritative only.", "Overlap"),
    "Confidence": ("Quality and review", "Confidence in the OS/OSM match: High (overlap), Medium, Low.", "High"),
    "Review Reason": ("Quality and review", "Why the record was flagged for a human look.", ""),
    "Review Status": ("Quality and review", "Blank = nothing to review; Open = needs a decision; Decided = decision recorded in overrides.csv.", "Open"),
    "Data Quality Score": ("Quality and review", "0-100 score of how well the record is evidenced (see 'Quality score' sheet).", "82"),
    "Data Quality Band": ("Quality and review", f"Band of the score: {BANDS}.", "High"),
    # location
    "Latitude": ("Location", "Latitude (WGS84, EPSG:4326) of a point inside the record.", "52.2406"),
    "Longitude": ("Location", "Longitude (WGS84, EPSG:4326) of a point inside the record.", "0.7167"),
    "Geom": ("Location", "Geometry as WKT in British National Grid (EPSG:27700). The GeoPackage stores the real geometry instead.", "POLYGON((...))"),
    "Count UPRNs": ("Location", "RESERVED: number of UPRNs inside the footprint. Not yet populated.", ""),
    "UPRNs inside footprint": ("Location", "RESERVED: list of UPRNs inside the footprint. Not yet populated.", ""),
    # audit
    "Created Date": ("Audit", "Date the record (or site) first appeared.", "2026-10-01"),
    "Created By": ("Audit", "Who created it ('pipeline' for automatic).", "pipeline"),
    "Last Edit Date": ("Audit", "Date of the last run that produced the record.", "2026-10-01"),
    "Last Edit By": ("Audit", "Who last edited it.", "pipeline"),
    "Rule Version": ("Audit", "Version stamp of the rules workbook used (date + content hash).", "2026-10-01-9d78ab"),
    "Run Date": ("Audit", "Date of the pipeline run.", "2026-10-01"),
    # reserved
    "DESNZ Priority Consumer Category A Default": ("Reserved", "RESERVED, not populated. Kept so the schema matches the original design.", ""),
    "BUUK Catastrophic Risk - Public Default": ("Reserved", "RESERVED, not populated.", ""),
    "BUUK Catastrophic Risk - Environmental Default": ("Reserved", "RESERVED, not populated.", ""),
    "BUUK Catastrophic Risk - Reputational Default": ("Reserved", "RESERVED, not populated.", ""),
    "DESNZ Priority Consumer Category A Override": ("Reserved", "RESERVED, not populated.", ""),
    "BUUK Catastrophic Risk - Public Override": ("Reserved", "RESERVED, not populated.", ""),
    "BUUK Catastrophic Risk - Environmental Override": ("Reserved", "RESERVED, not populated.", ""),
    "BUUK Catastrophic Risk - Reputational Override": ("Reserved", "RESERVED, not populated.", ""),
    "BUUK Catastrophic Risk - Catastrophic Risk Override": ("Reserved", "RESERVED, not populated.", ""),
}
# columns that exist only on the sites file, or mean something different there
SITE_ONLY = {
    "Primary Buildings": ("Site", "Number of buildings at the site whose role is Primary, Standalone or Whole site.", "3"),
    "Primary Category": ("Classification", "Category of the site.", "Education"),
    "Operational Importance": ("Classification", "Highest importance among the site's buildings.", "High"),
    "Site Method": ("Site", "How the site was formed (see buildings file).", "OS Functional Site"),
    "Created Date": ("Audit", "Date the site first appeared.", "2026-10-01"),
}

SOURCES = [
    ("OpenStreetMap", "Tagged features (amenities, healthcare, emergency, etc.) via the Overpass API",
     "Open Database Licence (ODbL) 1.0", "(c) OpenStreetMap contributors", "osm_features.json; osm_area_buildings.json; suffolk_boundary.json"),
    ("OS Open Map Local", "ImportantBuilding, FunctionalSite, RailwayStation and Building layers (GB grid squares TL, TM, TG)",
     "Open Government Licence v3.0", "Contains OS data (c) Crown copyright and database right 2026", "opmplc_essh_tl.zip; opmplc_essh_tm.zip; opmplc_essh_tg.zip"),
    ("OS Code-Point Open", "Postcode to coordinates (NHS ODS sites)", "Open Government Licence v3.0",
     "Contains OS data (c) Crown copyright and database right 2026; contains Royal Mail data (c) Royal Mail copyright and database right 2026; contains National Statistics data (c) Crown copyright and database right 2026",
     "codepo_gb.zip"),
    ("DfE Get Information About Schools", "Open schools, phase, pupil numbers, location", "Open Government Licence v3.0",
     "Contains public sector information licensed under the Open Government Licence v3.0. Source: Department for Education", "gias_edubasealldata.csv"),
    ("CQC HSCA active locations", "Care homes, hospices, hospitals, GPs, clinics: beds, rating, UPRN, location", "Open Government Licence v3.0",
     "Contains public sector information licensed under the Open Government Licence v3.0. Source: Care Quality Commission", "cqc_active_locations.ods"),
    ("NHS Organisation Data Service", "Trust sites, trusts and GP practices: ODS codes, status, parent trust", "Open Government Licence v3.0",
     "Contains public sector information licensed under the Open Government Licence v3.0. Source: NHS England Organisation Data Service", "ods_ets.csv; ods_etr.csv; ods_epraccur.csv"),
]

LIMITATIONS = [
    "Pilot area only: Suffolk. Counts and coverage will differ elsewhere.",
    "Importance comes from the Rules sheet by building type. It is not an assessment of each individual building.",
    "OS Open Map Local has no coastguard, lifeboat, government or conference classes, so those types come from OSM only and depend on OSM mapping.",
    "OSM-only records are as good as the volunteer mapping. Verification only exists for types a register covers (schools, care homes, hospitals, hospices, GPs, clinics, mental health).",
    "Private nurseries are not in the DfE data, so nursery sites show Not covered rather than Unverified.",
    "'Closed per register' means the registered organisation is closed; the building may still be in use by another organisation.",
    "Building Role is a footprint-size and OSM-tag heuristic (default 150 m2). A small but critical building can be labelled Ancillary; use overrides.",
    "Name matching is heuristic. Matches more than 500 m from the register's point are flagged for review; exact-name matches are accepted up to 2 km.",
    "Hospital sites collect every CQC service and NHS ODS code at the campus, so Authoritative Name/Rating can list several values.",
    "UPRN counts (Count UPRNs, UPRNs inside footprint) are not populated yet. UPRN holds only the UPRN(s) CQC gives for care locations.",
    "DESNZ and BUUK columns are reserved and empty.",
    "The data quality score is a heuristic for sorting and filtering, not a probability that a record is correct.",
    "OSM data is licensed under ODbL (share-alike). Check with your information governance or legal team before sharing this dataset outside the organisation.",
]


def _check_complete(df, sites):
    missing = [c for c in list(df.columns) if c not in D]
    missing += [c for c in sites.columns if c != "geometry" and c not in D and c not in SITE_ONLY]
    if missing:
        raise SystemExit("Data dictionary is missing descriptions for: " + ", ".join(sorted(set(missing))))


def _source_dates(root):
    out = []
    for name, use, lic, attr, files in SOURCES:
        dates, size = [], 0
        for f in [x.strip() for x in files.split(";")]:
            p = Path(root) / "data" / f
            if p.exists():
                dates.append(dt.date.fromtimestamp(p.stat().st_mtime))
                size += p.stat().st_size
        out.append((name, use, lic, attr, files, str(max(dates)) if dates else "", round(size / 1e6, 1)))
    return out


def write_docs(df, sites, version, today, changes, out, root):
    _check_complete(df, sites)
    df = df.copy()

    # ---- slim view ----
    slim_cols = ["Building ID", "Name", "Building Type", "Primary Category", "Operational Importance", "Building Role",
                 "Site ID", "Site Name", "Verification", "Data Quality Band", "Data Quality Score", "Source",
                 "Authoritative Source", "Capacity", "Capacity Unit", "Latitude", "Longitude"]
    df[slim_cols].to_csv(out / "critical_buildings_summary.csv", index=False)

    # ---- attribution ----
    src = _source_dates(root)
    lines = ["ATTRIBUTION AND LICENCES", "",
             "This dataset combines the sources below. Include the attribution lines when you share or publish it.", ""]
    for name, use, lic, attr, files, date, size in src:
        lines += [f"{name} - {lic}", f"  {attr}", ""]
    lines += ["Note: OpenStreetMap data is under the ODbL (share-alike). Check with your information governance "
              "or legal team before distributing this dataset outside the organisation.", ""]
    (out / "ATTRIBUTION.txt").write_text("\n".join(lines), encoding="utf-8")

    # ---- workbook ----
    def dict_rows(cols, extra):
        rows = []
        for c in cols:
            if c == "geometry":
                continue
            g, desc, ex = extra.get(c) or D[c]
            rows.append((g, c, desc, ex))
        return pd.DataFrame(rows, columns=["Group", "Column", "What it means", "Allowed values / example"])

    b_cols = list(df.columns)
    s_cols = list(sites.columns)
    readme = pd.DataFrame({"Critical Buildings dataset - read me": [
        f"Run date: {today}   Rules version: {version}",
        "Area: Suffolk (pilot).",
        "Files: critical_buildings.gpkg / .csv (one row per building), critical_sites.gpkg / .csv (one row per site), "
        "critical_buildings_summary.csv (slim view), run_summary.xlsx (QA), qa_map.html (map).",
        "Start with the summary CSV. Use 'Building Type', 'Operational Importance', 'Building Role' and 'Verification'.",
        "Fields ending in Default come from the rules; fields ending in Override are manual. 'Building Type' and similar "
        "un-suffixed fields show the final value.",
        "Sheets: Buildings columns, Sites columns, Values, Quality score, Sources and licences, Methodology, Known limitations, Release.",
    ]})
    values = pd.DataFrame([
        ("Source", "OSM", "Found in OpenStreetMap only"), ("Source", "OS Important Buildings", "Found in OS Open Map Local only"),
        ("Source", "Both", "OSM and OS agree it is the same building"),
        ("Source", "DfE GIAS / CQC / NHS ODS", "Added from an authoritative register that had no matching building"),
        ("Building Role", "Standalone", "Only building at its site"), ("Building Role", "Primary", "Main building at a multi-building site"),
        ("Building Role", "Ancillary", "Outbuilding, shed, garage or small building"),
        ("Building Role", "Whole site", "The record is the site area itself (no buildings found inside)"),
        ("Building Role", "Unknown", "Point only, no footprint, at a multi-building site"),
        ("Verification", "Verified", "A register matched this site"), ("Verification", "Verified (manual)", "A reviewer confirmed it against a register"),
        ("Verification", "Authoritative only", "Added from a register; OSM and OS did not have it"),
        ("Verification", "Unverified", "A register covers this type but nothing matched; may be closed, mis-tagged or unregistered"),
        ("Verification", "Closed per register", "The matching register record is closed"),
        ("Verification", "Not covered", "No register exists for this building type"),
        ("Confidence", "High / Medium / Low", "High = OS and OSM footprints overlap; Low = only close together with differing names"),
        ("Operational Importance", "High / Medium / Low", "From the Rules sheet by building type"),
    ], columns=["Column", "Value", "Meaning"])
    quality = pd.DataFrame([
        ("Authoritative check (0-40)", "Verified, Verified (manual), Authoritative only = 40; Not covered = 20; Unverified = 10; Closed per register = 0"),
        ("Source agreement (0-25)", "Both OSM and OS = 25; register only = 20; OSM only or OS only = 12"),
        ("Match confidence (0-20)", "High = 20; Medium = 12; Low = 5"),
        ("Geometry (0-15)", "Building footprint polygon = 15; point only = 5"),
        ("Open review item", "-10 while the record has an Open review status"),
        ("Total", f"Sum, capped 0-100. Bands: {BANDS}. A heuristic for sorting and filtering, not a probability."),
    ], columns=["Component", "Points"])
    sources = pd.DataFrame(src, columns=["Source", "What we use", "Licence", "Attribution", "Files", "Downloaded", "Size (MB)"])
    method = pd.DataFrame({"Step": [
        "1. Rules check: the Rules, Role settings and Reference mapping sheets are validated; the run stops with plain-English errors if wrong.",
        "2. Classify: each OSM feature (tags) and OS Important Building / station (theme > classification) is matched to a rule. The lowest Priority number wins.",
        "3. Merge: an OS building and an OSM feature are the same building if same category (same building type outside Education) and they overlap "
        f"(>= {merge.OVERLAP_FRACTION:.0%} of the OS polygon) or are within {merge.POINT_DISTANCE_M} m "
        f"(stations {merge.POINT_DISTANCE_BY_TYPE.get('Railway Station')} m). Proximity-only matches need the same type and similar names.",
        f"4. Sites: buildings are grouped by OS Functional Site, then OSM site/area polygon, then {sitelib.CLUSTER_GAP_M} m proximity cluster; "
        "same-named strays within "
        f"{sitelib.SITE_ATTACH_M} m join the site.",
        "5. Authoritative check: DfE, CQC and NHS ODS records are matched to sites by type, distance and name. Unmatched register records are added as new buildings; "
        "same-name records of a different type reclassify the site; sites no register recognises are marked Unverified.",
        "6. Roles: buildings are labelled Primary / Ancillary / Standalone / Whole site / Unknown from footprint size and OSM building tags (Role settings sheet).",
        "7. Overrides: manual decisions in outputs/overrides.csv are applied last and survive reruns. Building IDs and Site IDs stay stable between runs.",
        "8. Outputs, QA workbook, map and this dictionary are regenerated; changes since the previous run are listed in run_summary.xlsx.",
    ]})
    limits = pd.DataFrame({"Known limitation": LIMITATIONS})
    release = pd.DataFrame([
        ("Run date", str(today)), ("Rules version", version), ("Area", "Suffolk (pilot)"),
        ("Buildings", len(df)), ("Sites", len(sites)),
        ("Data quality bands", ", ".join(f"{k} {v}" for k, v in df["Data Quality Band"].value_counts().items())),
        ("Verification", ", ".join(f"{k} {v}" for k, v in df["Verification"].value_counts().items())),
        ("Review items open", int((df["Review Status"] == "Open").sum())),
        ("Changes since previous run", "first run" if changes is None else (
            ", ".join(f"{k} {v}" for k, v in changes.groupby("Change").size().items()) or "none")),
    ], columns=["Item", "Value"])

    path = out / "critical_buildings_data_dictionary.xlsx"
    with pd.ExcelWriter(path) as xw:
        readme.to_excel(xw, sheet_name="Read me", index=False)
        dict_rows(b_cols, {}).to_excel(xw, sheet_name="Buildings columns", index=False)
        dict_rows(s_cols, SITE_ONLY).to_excel(xw, sheet_name="Sites columns", index=False)
        values.to_excel(xw, sheet_name="Values", index=False)
        quality.to_excel(xw, sheet_name="Quality score", index=False)
        sources.to_excel(xw, sheet_name="Sources and licences", index=False)
        method.to_excel(xw, sheet_name="Methodology", index=False)
        limits.to_excel(xw, sheet_name="Known limitations", index=False)
        release.to_excel(xw, sheet_name="Release", index=False)
        head = PatternFill("solid", fgColor="1F3864")
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = head
            ws.freeze_panes = "A2"
            for col in ws.columns:
                width = min(90, max(14, max(len(str(c.value or "")) for c in col[:60]) + 2))
                ws.column_dimensions[col[0].column_letter].width = width
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
