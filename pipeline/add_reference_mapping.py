"""One-off: add the 'Reference mapping' sheet to the rules workbook.

Each row says: for reference source X, if field Y contains/equals Z, the record is Building Type T.
Rows are checked top to bottom within a source; the first match wins. Building Type IGNORE drops the record.
"""
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from rules_io import RULES_FILE

ROWS = [
    # source, field, match, value, building type, active, notes
    ("DfE", "EstablishmentStatus (name)", "Contains", "closed", "IGNORE", "Y", "Closed schools are dropped ('Open, but proposed to close' is kept)."),
    ("DfE", "EstablishmentStatus (name)", "Contains", "proposed to open", "IGNORE", "Y", "Not open yet."),
    ("DfE", "TypeOfEstablishment (name)", "Contains", "special", "School For Special Needs", "Y", "Community, academy, free and independent special schools."),
    ("DfE", "TypeOfEstablishment (name)", "Contains", "alternative provision", "School", "Y", "Alternative provision academies - generic school."),
    ("DfE", "TypeOfEstablishment (name)", "Contains", "pupil referral", "School", "Y", "Pupil referral units - generic school."),
    ("DfE", "TypeOfEstablishment (name)", "Contains", "higher education", "Further, Higher & University Education", "Y", ""),
    ("DfE", "TypeOfEstablishment (name)", "Contains", "further education", "Further, Higher & University Education", "Y", ""),
    ("DfE", "PhaseOfEducation (name)", "Contains", "16 plus", "Further, Higher & University Education", "Y", "Sixth-form colleges and 16-19 academies."),
    ("DfE", "PhaseOfEducation (name)", "Contains", "nursery", "Nursery & Child Care", "Y", "Maintained nursery schools only; private nurseries are not in this dataset."),
    ("DfE", "PhaseOfEducation (name)", "Contains", "primary", "Primary School", "Y", "Includes 'Middle deemed primary'."),
    ("DfE", "PhaseOfEducation (name)", "Contains", "secondary", "Secondary School", "Y", "Includes 'Middle deemed secondary'."),
    ("DfE", "PhaseOfEducation (name)", "Contains", "all-through", "Secondary School", "Y", "All-through schools cover primary and secondary; counted as secondary."),
    ("DfE", "PhaseOfEducation (name)", "Contains", "not applicable", "School", "Y", "Mostly independent schools with no phase - generic school."),
    ("CQC", "Dormant (Y/N)", "Equals", "Y", "IGNORE", "Y", "Dormant locations are dropped."),
    ("CQC", "Care home?", "Equals", "Y", "Care Home", "Y", "CQC flags residential care homes."),
    ("CQC", "Location Primary Inspection Category", "Contains", "hospice", "Hospice", "Y", ""),
    ("CQC", "Location Primary Inspection Category", "Contains", "acute hospital", "Hospital", "Y", "NHS and independent acute hospitals."),
    ("CQC", "Location Primary Inspection Category", "Contains", "mental health", "Mental health facilities", "Y", ""),
    ("CQC", "Location Primary Inspection Category", "Contains", "GP practices", "GP Surgery", "Y", ""),
    ("CQC", "Location Primary Inspection Category", "Contains", "independent consulting doctors", "Clinics offering Medical Treatment", "Y", ""),
    ("CQC", "Location Primary Inspection Category", "Contains", "community health", "Clinics offering Medical Treatment", "Y", ""),
    ("CQC", "Location Primary Inspection Category", "Contains", "ambulance", "Ambulance Station", "N", "Switched off: CQC ambulance locations are mostly independent patient transport, not emergency stations."),
    ("CQC", "Location Primary Inspection Category", "Contains", "community based adult social care", "IGNORE", "Y", "Home-care and supported-living agencies, mostly offices rather than care buildings."),
    ("CQC", "Location Primary Inspection Category", "Contains", "dentists", "IGNORE", "Y", "No dentist rule."),
    ("NHS ODS", "Name", "Contains", "hospital", "Hospital", "Y", "NHS trust sites with 'hospital' in the name (located by postcode)."),
]
wb = openpyxl.load_workbook(RULES_FILE)
if "Reference mapping" in wb.sheetnames:
    del wb["Reference mapping"]
ws = wb.create_sheet("Reference mapping", 3)
ws.append(["Reference Source", "Field", "Match", "Value", "Building Type", "Active (Y/N)", "Notes"])
for r in ROWS:
    ws.append(list(r))
head = PatternFill("solid", fgColor="1F3864")
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = head
for col, w in zip("ABCDEFG", [16, 38, 10, 34, 40, 12, 90]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
for formula, col in (('"DfE,CQC,NHS ODS"', "A"), ('"Contains,Equals"', "C"), ('"Y,N"', "F")):
    dv = DataValidation(type="list", formula1=formula, allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}500")

h = wb["How to edit"]
for text, style in [
    ("", None),
    ("Reference mapping (authoritative datasets)", "h"),
    ("DfE (schools), CQC (care and health) and NHS ODS records are matched to your Building Types using the 'Reference mapping' sheet.", None),
    ("Rows are checked top to bottom within each source; the first matching row wins. Building Type IGNORE drops the record.", None),
    ("Building Type must be spelled exactly as on the Rules sheet. The rules check flags any that are not.", None),
]:
    h.append([text])
    if style == "h":
        h.cell(row=h.max_row, column=1).font = Font(bold=True, size=12)
wb["Change log"].append(["Reference mapping", "New sheet mapping DfE / CQC / NHS ODS categories to Building Types"])
wb.save(RULES_FILE)
print(len(ROWS), "mapping rows added")
