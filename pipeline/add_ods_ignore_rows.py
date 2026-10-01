"""One-off: add NHS ODS IGNORE rows to the 'Reference mapping' sheet (before the generic 'hospital' row)."""
import openpyxl

from rules_io import RULES_FILE

wb = openpyxl.load_workbook(RULES_FILE)
ws = wb["Reference mapping"]
last = ws.max_row
assert ws.cell(last, 1).value == "NHS ODS" and ws.cell(last, 4).value == "hospital"
ws.insert_rows(last, 2)
rows = [("response post", "A response post is not a hospital."),
        ("road site", "'Hospital Road Site' style names are addresses, not hospitals.")]
for off, (value, note) in enumerate(rows):
    for col, val in enumerate(["NHS ODS", "Name", "Contains", value, "IGNORE", "Y", note], start=1):
        ws.cell(last + off, col).value = val
wb["Change log"].append(["Reference mapping", "Added NHS ODS IGNORE rows for 'response post' and 'road site' names"])
wb.save(RULES_FILE)
print("added", len(rows), "rows")
