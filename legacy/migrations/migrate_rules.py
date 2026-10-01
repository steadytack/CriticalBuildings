"""One-off migration: turn the original Rules sheet into the tidy, easy-to-edit format.

Reads  : Critcal Building Data Product.xlsx          (original, left untouched)
Writes : Critical Building Data Product (tidied).xlsx
         Sheets: How to edit | Rules | Output | Change log

After this runs, people edit the 'Rules' sheet in Excel. This script is not run again.
"""
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "Critcal Building Data Product.xlsx"
DST = ROOT / "Critical Building Data Product (tidied).xlsx"

changes = []  # (rule ref, what changed)

TEXT_FIXES = {
    "Unviersity": "University",
    "msuic": "music",
    "amenit=": "amenity=",
    "Nursey": "Nursery",
    "Dentention": "Detention",
}
# Rules whose generic matching should lose to more specific ones.
LOW_PRIORITY = {20: 90}   # generic "School"
DEFAULT_PRIORITY = 50
NONE_WORDS = ("none", "not listed")


def fix_text(ref, text):
    out = text
    for bad, good in TEXT_FIXES.items():
        if bad in out:
            out = out.replace(bad, good)
            changes.append((ref, f"Typo fixed: '{bad}' -> '{good}'"))
    return out


def split_osm(ref, raw):
    if raw is None or str(raw).strip().lower() in ("none", ""):
        return []
    raw = fix_text(ref, str(raw))
    # rule 12 used ':' between two tags instead of ';'
    if re.search(r"nursing_home:\s*building", raw):
        raw = re.sub(r"nursing_home:\s*building", "nursing_home; building", raw)
        changes.append((ref, "Split two tags that were joined by ':' into separate rows"))
    tags = []
    for part in raw.split(";"):
        part = part.strip()
        if not part:
            continue
        tidy = re.sub(r"\s*=\s*", "=", part)
        if tidy != part:
            changes.append((ref, f"Spacing tidied: '{part}' -> '{tidy}'"))
        if "=" not in tidy:
            # e.g. plant:output:electricity (no value) -> any value
            tidy += "=*"
            changes.append((ref, f"Tag had no value, set to wildcard: '{tidy}'"))
        tags.append(tidy)
    return tags


def split_os(ref, raw):
    if raw is None:
        return []
    raw = fix_text(ref, str(raw))
    if raw.strip().lower().startswith(NONE_WORDS):
        return []
    if ">" not in raw:
        return []
    group, classes = raw.split(">", 1)
    group = group.strip()
    return [f"{group} > {c.strip()}" for c in classes.split(";") if c.strip()]


wb = openpyxl.load_workbook(SRC)
src = wb["Rules"]
rules = []
for row in src.iter_rows(min_row=5, values_only=True):
    ref, cat, btype, imp, driver, os_tag, osm_tag = row[1:8]
    if ref is None:
        continue
    ref = int(ref)
    btype = fix_text(ref, str(btype)).strip()
    if str(btype) != str(row[3]):
        pass
    rules.append(dict(ref=ref, cat=str(cat).strip(), btype=btype, imp=str(imp).strip(),
                      driver=str(driver).strip(), os=split_os(ref, os_tag),
                      osm=split_osm(ref, osm_tag)))

# ---- build new workbook (keeps the Output sheet) ----
del wb["Rules"]
out = wb["Output"]
out["A1"] = ("Example structure only: rows show how individual buildings/features would be classified "
             "from OSM and/or OS Important Buildings using the rules on the 'Rules' sheet.")
changes.append(("Output sheet", "Header note said rules were in 'Sheet2'; now says 'Rules'"))

ws = wb.create_sheet("Rules", 0)
headers = ["Rule Ref", "Active (Y/N)", "Priority", "Primary Category", "Building Type",
           "Operational Importance", "Operational Importance Driver", "Source", "Tag", "Notes"]
ws.append(headers)
head_fill = PatternFill("solid", fgColor="1F3864")
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = head_fill
    c.alignment = Alignment(wrap_text=True, vertical="center")

alt = PatternFill("solid", fgColor="EEF3FA")
for i, r in enumerate(rules):
    prio = LOW_PRIORITY.get(r["ref"], DEFAULT_PRIORITY)
    rows = [("OS", t) for t in r["os"]] + [("OSM", t) for t in r["osm"]]
    notes = ""
    if not rows:
        rows = [("MANUAL", "")]
        notes = "No automatic tags: applied by manual override only."
    if r["ref"] == 20:
        notes = "Generic school rule - lower priority so specific school rules win."
    if r["ref"] == 22:
        notes = "landuse=port is an area: buildings inside the area are selected."
    if r["ref"] in (25, 27):
        notes = "Area tag: buildings inside the area are selected."
    for source, tag in rows:
        ws.append([r["ref"], "Y", prio, r["cat"], r["btype"], r["imp"], r["driver"], source, tag, notes])
        if i % 2:
            for c in ws[ws.max_row]:
                c.fill = alt
changes.append(("All rules", "One tag per row (was several tags in one cell separated by ';')"))
changes.append(("All rules", "Added Active (Y/N) and Priority columns (lower number wins when a feature matches several rules)"))

for col, w in zip("ABCDEFGHIJ", [9, 11, 9, 20, 40, 14, 55, 9, 55, 55]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
dv1 = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
dv2 = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=False)
dv3 = DataValidation(type="list", formula1='"OSM,OS,MANUAL"', allow_blank=False)
for dv, col in ((dv1, "B"), (dv2, "F"), (dv3, "H")):
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}1000")

# ---- how to edit ----
h = wb.create_sheet("How to edit", 0)
lines = [
    ("How to edit the rules", "title"),
    ("The 'Rules' sheet is the only place rules live. Change it here; no code changes are ever needed.", None),
    ("", None),
    ("Each ROW is one way of recognising a building type. A building type with three OSM tags has three rows.", None),
    ("", None),
    ("Columns", "h"),
    ("Rule Ref - number that groups the rows of one rule. Reuse it when adding another tag to an existing rule.", None),
    ("Active (Y/N) - set to N to switch a row off without deleting it.", None),
    ("Priority - if a feature matches several rules, the LOWEST number wins (10 beats 50). Normal = 50, generic = 90.", None),
    ("Primary Category / Building Type - labels that appear in the output dataset.", None),
    ("Operational Importance - High, Medium or Low (pick from the drop-down).", None),
    ("Operational Importance Driver - the plain-English reason for the importance.", None),
    ("Source - OSM (OpenStreetMap), OS (OS Important Buildings) or MANUAL (no automatic tag; set by hand).", None),
    ("Tag - what to look for. Formats below.", None),
    ("", None),
    ("Tag formats", "h"),
    ("OSM:  key=value, no spaces.          e.g.  amenity=fire_station", None),
    ("OSM any value:  key=*                e.g.  military_service=*", None),
    ("OS:   Group > Classification         e.g.  Emergency services > Fire Station", None),
    ("", None),
    ("Common tasks", "h"),
    ("Add a new tag to an existing rule: copy any row of that rule, keep the Rule Ref, change Source and Tag.", None),
    ("Add a new building type: use the next unused Rule Ref and fill every column on its row(s).", None),
    ("Change importance: change Operational Importance on ALL rows of that Rule Ref.", None),
    ("Stop using a rule: set Active to N on its rows.", None),
    ("", None),
    ("After editing", "h"),
    ("Save the file and run the pipeline. The first step is a rules check that explains any mistakes in plain English,", None),
    ("with the Rule Ref and row to fix. Nothing runs until the check passes. The output Rule Version is taken from this file.", None),
]
for text, style in lines:
    h.append([text])
    c = h.cell(row=h.max_row, column=1)
    if style == "title":
        c.font = Font(bold=True, size=16)
    elif style == "h":
        c.font = Font(bold=True, size=12)
h.column_dimensions["A"].width = 130

# ---- change log ----
cl = wb.create_sheet("Change log")
cl.append(["Rule Ref", "Change made during tidy-up"])
for c in cl[1]:
    c.font = Font(bold=True)
seen = set()
for ref, msg in changes:
    if (ref, msg) not in seen:
        seen.add((ref, msg))
        cl.append([ref, msg])
cl.column_dimensions["A"].width = 14
cl.column_dimensions["B"].width = 110

wb.save(DST)
print(f"Saved {DST.name}: {len(rules)} rules, {ws.max_row - 1} rule rows, {len(seen)} logged changes")
