"""One-time import of the legacy Excel rules workbook (and the CSV state files) into the control store."""
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pandas as pd

from critical_buildings.control import store as st

MIGRATION_USER = "migration"


def _s(v):
    return "" if v is None else str(v).strip()


def read_workbook(path):
    """Parse the tidied rules workbook into plain dicts (rules, tags, role settings, reference mapping)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    rules, tags, problems = {}, [], []
    for n, r in enumerate(wb["Rules"].iter_rows(min_row=2, max_col=10, values_only=True), start=2):
        if not any(v not in (None, "") for v in r):
            continue
        ref, active, prio, cat, btype, imp, driver, source, tag, notes = r
        ref = int(ref)
        rule = dict(rule_ref=ref, primary_category=_s(cat), building_type=_s(btype), importance=_s(imp),
                    importance_driver=_s(driver), priority=int(prio), active=1, notes="")
        if ref in rules:
            for k in ("primary_category", "building_type", "importance", "importance_driver", "priority"):
                if rules[ref][k] != rule[k]:
                    problems.append(f"Rule {ref} row {n}: {k} differs from the rule's first row ('{rules[ref][k]}' vs '{rule[k]}')")
        else:
            rules[ref] = rule
        tags.append(dict(rule_ref=ref, source=_s(source).upper(), tag=_s(tag), active=1 if _s(active).upper() == "Y" else 0,
                         notes=_s(notes)))
    if problems:
        raise ValueError("Workbook rows of one rule disagree:\n  " + "\n  ".join(problems))

    role = {"min_area": 150.0, "ancillary": [], "primary": [], "per_type": {}}
    in_types = False
    for a, b, _c in wb["Role settings"].iter_rows(min_row=2, max_col=3, values_only=True):
        if a == "Building Type":
            in_types = True
            continue
        if a in (None, ""):
            continue
        if not in_types:
            if a.startswith("Minimum primary"):
                role["min_area"] = float(b)
            elif a.startswith("Ancillary"):
                role["ancillary"] = [v.strip() for v in str(b or "").split(";") if v.strip()]
            elif a.startswith("Primary"):
                role["primary"] = [v.strip() for v in str(b or "").split(";") if v.strip()]
        elif b not in (None, ""):
            role["per_type"][a] = float(b)

    mapping = []
    for n, r in enumerate(wb["Reference mapping"].iter_rows(min_row=2, max_col=7, values_only=True), start=1):
        if not any(r):
            continue
        src, field, match, value, btype, active, notes = [_s(x) for x in r[:7]]
        mapping.append(dict(sort_order=n, source=src, field=field, match=match, value=value, building_type=btype,
                            active=1 if active.upper() == "Y" else 0, notes=notes))
    return list(rules.values()), tags, role, mapping


def import_workbook(store, path, reason="Imported from the Excel rules workbook"):
    """Replace the draft with the workbook content (used once, on an empty store)."""
    rules, tags, role, mapping = read_workbook(path)
    store.apply_table("rule", rules, MIGRATION_USER, reason)
    store.apply_table("rule_tag", tags, MIGRATION_USER, reason)
    store.apply_table("reference_mapping", mapping, MIGRATION_USER, reason)
    store.set_role_settings(role["min_area"], role["ancillary"], role["primary"], role["per_type"], MIGRATION_USER, reason)
    return {"rules": len(rules), "tags": len(tags), "mapping": len(mapping)}


def import_overrides_csv(store, path, reason="Imported from overrides.csv"):
    path = Path(path)
    df = pd.read_csv(path, dtype=str).fillna("")
    inv = {v: k for k, v in st.OVERRIDE_CSV_NAMES.items()}
    when = datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat(timespec="seconds")
    rows = []
    for rec in df.to_dict("records"):
        row = {inv[c]: v for c, v in rec.items() if c in inv}
        row.update(decided_by=MIGRATION_USER, decided_at=when)
        rows.append(row)
    store.save_overrides(rows, MIGRATION_USER, reason)
    return len(rows)


def import_registries(store, id_registry_csv, site_registry_csv, run_id="migration"):
    ids = pd.read_csv(id_registry_csv, dtype=str)
    with store.conn() as c:
        for r in ids.itertuples(index=False):
            c.execute("INSERT OR REPLACE INTO building_registry (building_id, os_key, osm_key, created, wkt, first_seen_run, last_seen_run, retired_run) "
                      "VALUES (?,?,?,?,?,?,?,NULL)",
                      (r.building_id, r.os_key if isinstance(r.os_key, str) else None, r.osm_key if isinstance(r.osm_key, str) else None,
                       r.created, r.wkt, run_id, run_id))
        sites = pd.read_csv(site_registry_csv, dtype=str)
        for r in sites.itertuples(index=False):
            c.execute("INSERT OR REPLACE INTO site_registry (site_id, created, first_seen_run, last_seen_run, retired_run) VALUES (?,?,?,?,NULL)",
                      (r.site_id, r.created, run_id, run_id))
    return len(ids), len(sites)
