"""A read-only Excel view of a published rule set, for people who prefer to read rules in a spreadsheet.

The workbook is generated from the control store and is NOT a place to edit: changes made in it are ignored.
Edit rules in the UI (python -m critical_buildings ui).
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from critical_buildings import paths
from critical_buildings.control import ruleset as rs

HEAD = PatternFill("solid", fgColor="1F3864")
NOTE = PatternFill("solid", fgColor="FFF2CC")


def _sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(list(r))
    for col, w in zip("ABCDEFGHIJKL", widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def export(store, out_path=None, version_id=None):
    """Write the workbook; returns its path."""
    vid, snap = store.get_version(version_id)
    label = rs.version_label(vid, snap)
    created = store.versions_df().set_index("version_id").loc[vid, "created_at"]
    out_path = Path(out_path) if out_path else paths.CONTROL / "exports" / f"rules_{label}.xlsx"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Read me"
    for line in [f"Critical Buildings rules - version {label}", f"Published {created}", "",
                 "THIS WORKBOOK IS A READ-ONLY VIEW. Changes made here are ignored.",
                 "To change the rules use the editor:  python -m critical_buildings ui",
                 "", "Sheets: Rule summary (one row per rule), Rules (one row per tag), Role settings, Reference mapping."]:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=14)
    ws["A4"].fill = NOTE
    ws["A4"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 90

    rules = {r["rule_ref"]: r for r in snap["rules"]}
    tags_by_rule = {}
    for t in snap["tags"]:
        tags_by_rule.setdefault(t["rule_ref"], []).append(t)
    _sheet(wb, "Rule summary",
           ["Rule Ref", "Active", "Priority", "Primary Category", "Building Type", "Operational Importance",
            "Operational Importance Driver", "Active tags", "Notes"],
           [(r["rule_ref"], "Y" if r["active"] else "N", r["priority"], r["primary_category"], r["building_type"], r["importance"],
             r["importance_driver"], sum(1 for t in tags_by_rule.get(r["rule_ref"], []) if t["active"]), r["notes"])
            for r in snap["rules"]], [9, 8, 9, 20, 40, 14, 55, 12, 40])
    _sheet(wb, "Rules",
           ["Rule Ref", "Active (rule and tag)", "Priority", "Primary Category", "Building Type", "Operational Importance",
            "Operational Importance Driver", "Source", "Tag", "Notes"],
           [(t["rule_ref"], "Y" if (rules[t["rule_ref"]]["active"] and t["active"]) else "N", rules[t["rule_ref"]]["priority"],
             rules[t["rule_ref"]]["primary_category"], rules[t["rule_ref"]]["building_type"], rules[t["rule_ref"]]["importance"],
             rules[t["rule_ref"]]["importance_driver"], t["source"], t["tag"], t["notes"]) for t in snap["tags"]],
           [9, 12, 9, 20, 40, 14, 55, 9, 55, 50])
    role = snap["role"]
    rows = [("Minimum primary footprint (m2)", role["min_area"], "At least this big = Primary; smaller = Ancillary."),
            ("Ancillary OSM building values", "; ".join(role["ancillary"]), "Ancillary whatever the size."),
            ("Primary OSM building values", "; ".join(role["primary"]), "Primary whatever the size.")]
    rows += [(f"Minimum for {bt}", v, "Per-type override of the minimum.") for bt, v in sorted(role["per_type"].items())]
    ws = _sheet(wb, "Role settings", ["Setting", "Value", "Explanation"], rows, [40, 90, 60])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    _sheet(wb, "Reference mapping", ["Order", "Reference Source", "Field", "Match", "Value", "Building Type", "Active", "Notes"],
           [(m["sort_order"], m["source"], m["field"], m["match"], m["value"], m["building_type"], "Y" if m["active"] else "N", m["notes"])
            for m in snap["reference_mapping"]], [8, 16, 36, 10, 34, 40, 8, 80])
    wb.save(out_path)
    return out_path
