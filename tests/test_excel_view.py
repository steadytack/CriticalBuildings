import openpyxl

from critical_buildings.control import excel_view


def test_excel_view_is_a_faithful_read_only_copy_of_the_published_rules(seeded_store, tmp_path):
    seeded_store.publish("ann", "v1")
    out = excel_view.export(seeded_store, tmp_path / "rules.xlsx")
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Read me", "Rule summary", "Rules", "Role settings", "Reference mapping"]
    readme = [c.value for c in wb["Read me"]["A"] if c.value]
    assert any("READ-ONLY VIEW" in line for line in readme) and readme[0].startswith("Critical Buildings rules - version v1-")
    summary = list(wb["Rule summary"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in summary] == [1, 17, 20, 34] and summary[0][4] == "Fire Station" and summary[0][7] == 2
    tags = list(wb["Rules"].iter_rows(min_row=2, values_only=True))
    assert len(tags) == 6 and tags[0][8] == "amenity=fire_station" and tags[0][1] == "Y"
    role = {r[0]: r[1] for r in wb["Role settings"].iter_rows(min_row=2, values_only=True)}
    assert role["Minimum primary footprint (m2)"] == 150.0 and "garage" in role["Ancillary OSM building values"]
    assert wb["Reference mapping"].max_row == 3
